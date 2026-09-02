#!/usr/bin/env python3
"""
parse-xlsx.py —— 把学员汇总 Excel 解析为结构化 JSON（只读，不写任何数据）

职责单一：xlsx → JSON。写入云端由 import-students.mjs 负责。
这样拆分的好处：解析逻辑可单独调试，换表格格式只改这一个文件。

用法：
    python3 scripts/parse-xlsx.py <xlsx路径> [工作表名] [表头所在行号]

输出：JSON 到 stdout
    { "ok": true, "rows": [...], "skipped": [...], "stats": {...} }

约定（与 Excel 模板一致，改列名需同步改这里）：
    学号 | 姓名 | 身份职业 | AI使用经验 | 用过哪些AI | 学习目的 | 优先学习方向 | 补充开放题 | 备注
"""
import json
import sys

try:
    import openpyxl
except ImportError:
    print(json.dumps({"ok": False, "error": "缺少 openpyxl，请先安装：pip install openpyxl"}))
    sys.exit(1)

PLACEHOLDERS = {"待补", "—", "-", "", None}


def clean(v):
    """把占位符统一转成空字符串，便于后续判断。"""
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s in PLACEHOLDERS else s


def main():
    if len(sys.argv) < 2:
        print(json.dumps({"ok": False, "error": "用法: parse-xlsx.py <xlsx路径> [表名] [表头行号]"}))
        sys.exit(1)

    path = sys.argv[1]
    sheet_name = sys.argv[2] if len(sys.argv) > 2 else None
    header_row = int(sys.argv[3]) if len(sys.argv) > 3 else 3

    wb = openpyxl.load_workbook(path, data_only=True)

    # 未指定表名时，挑第一个含有「学号」列的工作表
    ws = None
    if sheet_name:
        ws = wb[sheet_name]
    else:
        for cand in wb.worksheets:
            for r in range(1, min(8, cand.max_row + 1)):
                vals = [str(c.value).strip() if c.value else "" for c in cand[r]]
                if "学号" in vals:
                    ws = cand
                    header_row = r
                    break
            if ws:
                break
    if ws is None:
        print(json.dumps({"ok": False, "error": "找不到含「学号」列的工作表"}))
        sys.exit(1)

    header = [str(c.value).strip() if c.value else "" for c in ws[header_row]]
    col = {name: i for i, name in enumerate(header) if name}

    def get(row, name):
        i = col.get(name)
        if i is None or i >= len(row):
            return ""
        return clean(row[i])

    rows, skipped = [], []
    for r in range(header_row + 1, ws.max_row + 1):
        raw = [c.value for c in ws[r]]
        # 完全空行 → 跳过
        if not any(x is not None and str(x).strip() for x in raw):
            continue

        sid_raw = raw[col["学号"]] if "学号" in col and col["学号"] < len(raw) else None
        sid = clean(sid_raw)
        name = get(raw, "姓名")

        # 学号非数字 → 视作分隔行/未署名问卷，单独收集供人工处理
        if not sid or not str(sid).strip().isdigit():
            label = str(sid_raw).strip() if sid_raw else f"第{r}行"
            if any(str(x).strip() for x in raw if x is not None):
                skipped.append({"row": r, "label": label, "name": name or "(未填)", "raw": [str(x) if x is not None else "" for x in raw]})
            continue

        rows.append({
            "seq": int(sid),
            "name": name,
            "occupation": get(raw, "身份职业"),
            "ai_experience": get(raw, "AI使用经验"),
            "ai_tools": get(raw, "用过哪些AI"),
            "goal": get(raw, "学习目的"),
            "priority": get(raw, "优先学习方向"),
            "open_answer": get(raw, "补充开放题"),
            "remark": get(raw, "备注"),
        })

    filled = sum(1 for x in rows if x["occupation"])
    print(json.dumps({
        "ok": True,
        "sheet": ws.title,
        "header_row": header_row,
        "columns": header,
        "rows": rows,
        "skipped": skipped,
        "stats": {"total": len(rows), "filled": filled, "pending": len(rows) - filled, "skipped": len(skipped)},
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
