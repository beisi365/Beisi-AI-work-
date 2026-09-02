#!/usr/bin/env python3
"""
build-xlsx.py —— 把学员档案 JSON 渲染成 Excel（供 export-students.mjs 调用）

职责单一：JSON → xlsx。数据拉取由 export-students.mjs 负责。

用法：
    python3 scripts/build-xlsx.py <payload.json> <输出路径> [班级id]

payload.json 结构：
    {
      "columns": [["列标题", "字段key"], ...],
      "editableCount": 9,          # 前 N 列可回写，后面是只读参考列
      "rows": [{"列标题": "值", ...}, ...]
    }

排版约定（务必与 parse-xlsx.py 的自动识别逻辑兼容）：
    - 主表「学员名单」放在第一个工作表，表头固定在第 1 行；
    - 说明文字一律放第二个工作表，避免开头出现含「学号」的干扰行，
      否则 parse-xlsx.py 会把说明行误判成表头。
"""
import json
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("缺少 openpyxl，请先安装：pip install openpyxl")
    sys.exit(1)


def main():
    if len(sys.argv) < 3:
        print("用法: build-xlsx.py <payload.json> <输出路径> [班级id]")
        sys.exit(1)

    src, out = sys.argv[1], sys.argv[2]
    class_id = sys.argv[3] if len(sys.argv) > 3 else ""

    with open(src, "r", encoding="utf-8") as f:
        payload = json.load(f)

    columns = payload["columns"]
    editable_count = payload.get("editableCount", len(columns))
    rows = payload.get("rows", [])

    wb = Workbook()

    # ---------- Sheet1：学员名单（表头必须在第 1 行） ----------
    ws = wb.active
    ws.title = "学员名单"

    header_fill = PatternFill("solid", fgColor="2F3A4A")
    editable_fill = PatternFill("solid", fgColor="FFF7E8")  # 可编辑列：暖底
    ref_fill = PatternFill("solid", fgColor="F2F4F7")        # 只读列：灰底
    header_font = Font(color="FFFFFF", bold=True, size=11)

    for c, (label, _key) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=c, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r, row in enumerate(rows, start=2):
        for c, (label, _key) in enumerate(columns, start=1):
            cell = ws.cell(row=r, column=c, value=row.get(label, ""))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if c <= editable_count:
                cell.fill = editable_fill

    # 只读参考列整列置灰，视觉上区分
    for c in range(editable_count + 1, len(columns) + 1):
        for r in range(1, len(rows) + 2):
            ws.cell(row=r, column=c).fill = ref_fill

    # 列宽：按内容长度自适应，限幅避免过长
    for c, (label, _key) in enumerate(columns, start=1):
        longest = max([len(str(label))] + [len(str(row.get(label, ""))) for row in rows] or [10])
        width = min(max(longest + 4, 10), 42)
        ws.column_dimensions[get_column_letter(c)].width = width

    ws.freeze_panes = "C2"  # 冻结表头与前两列（学号/姓名），横向滚动时不丢锚点

    # ---------- Sheet2：填写说明 ----------
    tip = wb.create_sheet("填写说明")
    tips = [
        ["学员档案 · 外部编辑说明", ""],
        ["", ""],
        ["导出班级", class_id or "全部"],
        ["导出时间", __import__("datetime").datetime.now().strftime("%Y-%m-%d %H:%M")],
        ["", ""],
        ["怎么用", ""],
        ["1. 编辑", "直接改「学员名单」里的单元格，可批量填充、下拉复制。"],
        ["2. 保存", "改完直接保存，不用另存为 CSV，保持 .xlsx 即可。"],
        ["3. 回导", "把文件放进项目 imports/ 目录，然后对助手说「导入学员表」。"],
        ["", ""],
        ["注意事项", ""],
        ["学号列不要改", "学号是平台与表格对齐的主键。改了会被当成新学员，产生重复档案。"],
        ["新增学员", "在末尾新增一行并填上学号，导入时会自动建档（不要插在中间）。"],
        ["清空单元格", "留空表示「不改动」：导入不会用空值覆盖平台已有内容。"],
        ["确实要清空", "请填一个占位内容再改，或让助手用 --overwrite-empty 导入。"],
        ["", ""],
        ["列说明", ""],
        ["暖色列（前若干列）", "可编辑并回写到平台。"],
        ["灰色列（后面）", "只导出供对照，导入时会被忽略，改了也不会写入平台。"],
    ]
    for r, (a, b) in enumerate(tips, start=1):
        tip.cell(row=r, column=1, value=a).font = Font(bold=True)
        tip.cell(row=r, column=2, value=b).alignment = Alignment(wrap_text=True, vertical="top")
    tip.column_dimensions["A"].width = 22
    tip.column_dimensions["B"].width = 68

    wb.save(out)
    print(f"[build-xlsx] 已生成：{out}（{len(rows)} 行 × {len(columns)} 列）")


if __name__ == "__main__":
    main()
